"""Tests for the AI-generated groundwater assessment report."""

import io

from fastapi.testclient import TestClient
from openpyxl import load_workbook


def test_assessment_report_json(client: TestClient, auth_user):
    resp = client.get(
        "/api/reports/assessment",
        params={
            "state": "Andhra Pradesh",
            "district": "East Godavari",
            "year_from": 2020,
            "year_to": 2030,
        },
        headers=auth_user,
    )
    assert resp.status_code == 200, resp.text
    r = resp.json()
    assert r["scope"]["district"] == "East Godavari"
    assert r["period"]["from"] == 2020
    assert r["period"]["to"] == 2030
    assert r["is_demo"] is True
    assert len(r["executive_summary"]) >= 2
    assert r["status"]["assessment_units"] > 0
    assert r["status"]["stage"] > 0
    assert r["status"]["category"] in ("Safe", "Semi-critical", "Critical", "Over-exploited")
    assert r["status"]["water_level"]["value"] is not None
    assert r["trend"]["series"]
    assert r["prediction"]["forecast"]
    assert r["prediction"]["end_year"] == 2030
    assert r["prediction"]["end_value"] is not None
    assert r["risk"]["level"] in ("Low", "Medium", "High", "Critical")
    assert r["map"]["features"]
    assert r["recommendations"]
    assert len(r["sources"]) >= 4


def test_assessment_report_all_scope(client: TestClient, auth_user):
    resp = client.get("/api/reports/assessment", headers=auth_user)
    assert resp.status_code == 200, resp.text
    assert resp.json()["scope"]["display"]
    assert resp.json()["map"]["features"]


def test_assessment_report_unknown_district_404(client: TestClient, auth_user):
    resp = client.get(
        "/api/reports/assessment",
        params={"state": "Andhra Pradesh", "district": "NoSuchDistrict"},
        headers=auth_user,
    )
    assert resp.status_code == 404


def test_assessment_report_pdf(client: TestClient, auth_user):
    resp = client.get(
        "/api/reports/assessment.pdf",
        params={"state": "Andhra Pradesh", "district": "East Godavari", "year_to": 2030},
        headers=auth_user,
    )
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"] == "application/pdf"
    assert resp.content.startswith(b"%PDF")
    assert len(resp.content) > 1000


def test_assessment_report_xlsx(client: TestClient, auth_user):
    resp = client.get(
        "/api/reports/assessment.xlsx",
        params={
            "state": "Andhra Pradesh",
            "district": "East Godavari",
            "year_from": 2020,
            "year_to": 2030,
        },
        headers=auth_user,
    )
    assert resp.status_code == 200, resp.text
    wb = load_workbook(io.BytesIO(resp.content), read_only=True)
    for sheet in ("Summary", "Status", "Trend", "Prediction", "Risk", "Recommendations", "Map Data", "Sources"):
        assert sheet in wb.sheetnames, f"missing sheet {sheet}"